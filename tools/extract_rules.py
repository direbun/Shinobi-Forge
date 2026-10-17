"""Extract the usable Naruto 5e rules catalogue from the supplied sheet export.

Usage:
  python tools/extract_rules.py source-character-sheet.xlsx data/rules.json [pdf ...]

The workbook is treated as untrusted input. This script only reads cell values and
does not evaluate formulas, follow links, or execute embedded content.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


LEVELS = list(range(1, 21))
CLAN_LEVELS = [1, 3, 5, 7, 11, 15, 18]

STARTING_EQUIPMENT = {
    "Genjutsu-Specialist": [
        ["1 Simple Weapon"], ["Kunai Stack", "Shuriken Stack"], ["Padded Cloth"], ["Poison Kit"], ["Smoke Bomb"],
    ],
    "Hunter-Nin": [
        ["Padded Cloth"], ["2 Simple Weapons", "1 Martial Weapon"], ["Kunai Stack", "Shuriken Stack"],
        ["Flash Tag", "Paper Bomb"], ["2 Kits of your choice"],
    ],
    "Intelligence-Operative": [
        ["Padded Cloth", "Combat Jacket"], ["1 Simple Weapon", "2 Knuckle Blades"], ["Kunai Stack", "Fuma Shuriken"],
        ["1 Tool Kit of your choice"], ["Paper Bomb"],
    ],
    "Medical-Nin": [
        ["Padded Cloth", "Combat Jacket"], ["1 Simple Weapon"], ["Kunai Stack", "Senbon Stack"], ["Flash Tag"], ["Medicine Kit"],
    ],
    "Ninjutsu-Specialist": [
        ["Padded Cloth"], ["1 Simple Weapon"], ["Kunai Stack", "Shuriken Stack"], ["D-Rank Ninjutsu Scroll"],
    ],
    "Scout-Nin": [
        ["Padded Cloth", "Combat Jacket", "Combat Armor"], ["1 Simple Weapon", "1 Martial Weapon"],
        ["Kunai Stack", "Shuriken Stack"], ["Paper Bomb", "Flash Tag"], ["1 Toolkit of your choice"],
    ],
    "Taijutsu-Specialist": [
        ["No Armor", "Padded Cloth", "Combat Jacket"], ["Combat Bracers", "Iron Claw"], ["Kunai Stack", "Shuriken Stack"],
        ["2 Paper Bombs", "2 Flash Tags"], ["Smoke Bomb"],
    ],
    "Weapon-Specialist": [
        ["Padded Cloth", "Combat Jacket", "Combat Armor"], ["1 Simple Weapon", "1 Martial Weapon"],
        ["1 Simple Weapon", "1 Martial Weapon"], ["Kunai Stack", "Shuriken Stack"], ["2 Paper Bombs"],
    ],
    "Puppet-Master": [
        ["Hand Crossbow and Bolt Stack", "2 Simple Weapons"], ["Kunai Stack", "Shuriken Stack"],
        ["Padded Cloth", "Combat Jacket"], ["Crafter's Pack"], ["Paper Bomb"],
    ],
    "Cooking-Nin": [
        ["Padded Cloth", "Combat Jacket"], ["1 Simple Weapon"], ["Kunai Stack", "Shuriken Stack"],
        ["Cooking Tools"], ["Flash Tag"], ["Paper Bomb"],
    ],
    "Science-Nin": [
        ["Padded Cloth", "Combat Jacket"], ["1 Simple Weapon"], ["Kunai Stack", "Shuriken Stack"],
        ["Paper Bomb", "Flash Tag"], ["2 Kits of your choice"],
    ],
    "Jujutsu-Sorcerer": [
        ["Padded Cloth", "Combat Jacket"], ["1 Martial Weapon", "1 Exotic Weapon"], ["Kunai Stack", "Chakram"],
        ["2 Simple Weapons"], ["2 Paper Bombs", "2 Flash Tags"], ["Tracking Kit", "Forensics Kit"],
    ],
    "Lunar-Slayer": [
        ["Padded Cloth", "Combat Jacket"], ["1 Simple Weapon", "1 Martial Weapon"], ["1 Simple Weapon", "1 Martial Weapon"],
        ["Shinobi Waist Bag, Belt Pouch, and Leg Pouch"], ["1 Kit of your choice"],
    ],
}

WEAPON_MARKET = {
    "kunai": (10, "Simple"), "hand axe": (10, "Simple"), "sai": (20, "Simple"),
    "tanto (shortsword)": (10, "Simple"), "kama (hand scythe)": (10, "Simple"),
    "gunsen": (10, "Simple"), "quarterstaff": (10, "Simple"), "spear": (10, "Simple"),
    "weighted chain": (10, "Simple"), "kusarigama (chained hand scythe)": (10, "Simple"),
    "tekko": (20, "Simple"), "senbon": (15, "Simple"), "short bow": (20, "Simple"),
    "shuriken": (10, "Simple"), "sling": (5, "Simple"), "light crossbow": (20, "Simple"),
    "bola": (5, "Simple"),
    "broadsword": (20, "Martial"), "iron claw": (20, "Martial"), "taichi": (40, "Martial"),
    "katana": (30, "Martial"), "odachi (great sword)": (80, "Martial"),
    "knuckle blades": (20, "Martial"), "hidden blade": (20, "Martial"),
    "chained spear": (20, "Martial"), "chigiriki": (25, "Martial"), "whip": (20, "Martial"),
    "battle wire": (20, "Martial"), "naginata": (40, "Martial"),
    "sasumata (forked spear)": (40, "Martial"), "great axe": (40, "Martial"),
    "scythe": (40, "Martial"), "tinbe rochin": (15, "Martial"), "yari": (40, "Martial"),
    "hooked lance": (40, "Martial"), "tetsubo": (40, "Martial"), "tonfa": (20, "Martial"),
    "war club": (40, "Martial"), "nunchaku": (20, "Martial"),
    "combat bracers": (20, "Martial"), "jitte": (20, "Martial"), "gunbai fan": (40, "Martial"),
    "kanabo": (30, "Martial"), "otsuchi hammer": (40, "Martial"), "chakram": (40, "Martial"),
    "monster chakram": (40, "Martial"), "fuma-shuriken": (20, "Martial"),
    "monster shuriken": (40, "Martial"), "torinawa": (5, "Martial"), "boomerang": (25, "Martial"),
    "monster boomerang": (50, "Martial"), "longbow": (40, "Martial"),
    "crossbow, hand": (20, "Martial"), "crossbow, heavy": (40, "Martial"),
    "blowgun": (20, "Martial"), "net": (20, "Martial"),
    "sansetsukon": (100, "Exotic"), "triple-bladed scythe": (100, "Exotic"),
    "cleaver sword": (100, "Exotic"), "triple katar": (100, "Exotic"), "urumi": (100, "Exotic"),
    "chokuto": (20, "Exotic"), "matchlock pistol": (200, "Exotic"),
    "matchlock rifle": (250, "Exotic"), "hiya taihou": (500, "Exotic"),
    "combat scroll": (250, "Exotic"), "giant combat scroll": (500, "Exotic"),
    "ballistic kunai gun": (300, "Exotic"), "ballistic shuriken rifle": (500, "Exotic"),
}

ARMOR_PRICES = {
    "padded cloth": 100, "leather weave": 250, "armored cloth": 500, "reinforced cloth": 750,
    "synthetic weave": 1500, "shinobi weave": 2500, "combat jacket": 100,
    "shinobi jacket": 250, "shinobi combat jacket": 500, "chunin jacket": 750,
    "battle coat": 1500, "armored chunin coat": 2500, "combat armor": 100,
    "synthetic armor": 250, "jonin armor": 500, "elite jonin armor": 750,
    "ronin armor": 1500, "samurai armor": 2500,
}

PROFICIENCY_DEGREES = [
    {"value": 0, "name": "Not Proficient", "rank": "Untrained", "masteryBonus": 0},
    {"value": 1, "name": "Proficient", "rank": "Genin", "masteryBonus": 0},
    {"value": 2, "name": "Mastery 1", "rank": "Chunin", "masteryBonus": 2},
    {"value": 3, "name": "Mastery 2", "rank": "Jonin", "masteryBonus": 4},
    {"value": 4, "name": "Mastery 3", "rank": "Sannin / Kage", "masteryBonus": 6},
]

ALL_SKILL_IDS = [
    "acrobatics", "animal-handling", "athletics", "chakra-control", "crafting", "deception",
    "history", "illusions", "insight", "intimidation", "investigation", "martial-arts",
    "medicine", "nature", "ninshou", "perception", "performance", "persuasion",
    "sleight-of-hand", "stealth", "survival",
]


def fixed_skills(source: str, *skills: str) -> dict:
    return {"id": f"{source}-fixed", "label": "Granted skills", "count": len(skills), "options": list(skills), "fixed": True}


def choose_skills(source: str, count: int, skills: list[str], label: str | None = None) -> dict:
    return {"id": f"{source}-choice", "label": label or f"Choose {count} skills", "count": count, "options": skills, "fixed": False}


CLASS_PROFICIENCY_GRANTS = {
    "Genjutsu-Specialist": [fixed_skills("class-genjutsu", "illusions"), choose_skills("class-genjutsu", 3, ["chakra-control", "deception", "history", "insight", "intimidation", "investigation", "perception", "persuasion", "stealth"])],
    "Hunter-Nin": [fixed_skills("class-hunter", "stealth"), choose_skills("class-hunter", 3, ALL_SKILL_IDS, "Choose three other skills")],
    "Intelligence-Operative": [fixed_skills("class-intelligence", "investigation"), choose_skills("class-intelligence", 3, ["chakra-control", "crafting", "ninshou", "persuasion", "illusions", "insight", "intimidation", "history", "perception", "martial-arts"])],
    "Medical-Nin": [fixed_skills("class-medical", "medicine", "chakra-control"), choose_skills("class-medical", 2, ["history", "insight", "perception", "illusions", "ninshou"])],
    "Ninjutsu-Specialist": [fixed_skills("class-ninjutsu", "ninshou"), choose_skills("class-ninjutsu", 4, ["martial-arts", "insight", "investigation", "history", "stealth", "perception", "chakra-control"])],
    "Scout-Nin": [choose_skills("class-scout", 3, ["ninshou", "illusions", "martial-arts", "insight", "investigation", "nature", "acrobatics", "athletics", "history", "stealth", "perception", "medicine", "chakra-control"])],
    "Taijutsu-Specialist": [fixed_skills("class-taijutsu", "martial-arts"), choose_skills("class-taijutsu", 2, ["ninshou", "illusions", "insight", "investigation", "nature", "acrobatics", "athletics", "history", "stealth", "perception", "medicine", "chakra-control"])],
    "Weapon-Specialist": [fixed_skills("class-weapon", "martial-arts"), choose_skills("class-weapon", 3, ["acrobatics", "athletics", "crafting", "history", "intimidation", "perception", "sleight-of-hand", "stealth", "survival"])],
    "Puppet-Master": [fixed_skills("class-puppet", "crafting"), choose_skills("class-puppet", 3, ["chakra-control", "deception", "history", "insight", "intimidation", "investigation", "ninshou", "perception", "nature", "sleight-of-hand", "stealth"])],
    "Cooking-Nin": [fixed_skills("class-cooking", "survival"), choose_skills("class-cooking", 3, ["athletics", "crafting", "deception", "history", "insight", "intimidation", "investigation", "medicine", "nature", "perception", "sleight-of-hand"])],
    "Science-Nin": [choose_skills("class-science", 4, ["ninshou", "insight", "investigation", "nature", "crafting", "athletics", "history", "stealth", "perception", "medicine", "chakra-control"])],
}

CLASS_COMBAT_PROFICIENCIES = {
    "Genjutsu-Specialist": {"armor": ["Light Armor"], "weapons": ["Simple Weapons"], "ninjaTools": ["Trapper's Kit"], "savingThrows": ["Constitution", "Charisma"]},
    "Hunter-Nin": {"armor": ["Light Armor"], "weapons": ["Simple Weapons", "Broadsword", "Iron Claw", "Katana", "Knuckle Blades", "Chakram"], "ninjaTools": ["Tracking Kit", "Forensics Kit"], "savingThrows": ["Dexterity", "Intelligence"]},
    "Intelligence-Operative": {"armor": ["Light Armor", "Medium Armor"], "weapons": ["Simple Weapons", "Fuma Shuriken", "Knuckle Blades"], "ninjaTools": ["Trapper's Kit", "Poison Kit", "Tracking Kit", "Forensics Kit", "Forgery Kit"], "savingThrows": ["Dexterity", "Intelligence"]},
    "Medical-Nin": {"armor": ["Light Armor", "Medium Armor"], "weapons": ["Simple Weapons"], "ninjaTools": ["Medicine Kit", "Poison Kit"], "savingThrows": ["Wisdom", "Charisma"]},
    "Ninjutsu-Specialist": {"armor": ["Light Armor"], "weapons": ["Simple Weapons"], "ninjaTools": [], "savingThrows": ["Wisdom", "Intelligence"]},
    "Scout-Nin": {"armor": ["Light Armor", "Medium Armor"], "weapons": ["Simple Weapons", "Martial Weapons"], "ninjaTools": ["Tracking Kit", "Trapper's Kit"], "savingThrows": ["Strength", "Constitution"]},
    "Taijutsu-Specialist": {"armor": [], "weapons": ["Simple Weapons", "Martial Weapons"], "ninjaTools": ["Tracking Kit", "Trapper's Kit"], "savingThrows": ["Strength", "Dexterity"]},
    "Weapon-Specialist": {"armor": ["Light Armor", "Medium Armor", "Heavy Armor"], "weapons": ["Simple Weapons", "Martial Weapons"], "ninjaTools": ["Trapper's Kit", "Poison Kit", "Tracking Kit", "Weaponsmith Kit", "Armorsmith Kit"], "savingThrows": ["Strength", "Constitution"]},
    "Puppet-Master": {"armor": ["Light Armor", "Medium Armor"], "weapons": ["Simple Weapons", "Battle Wire", "Hand Crossbow", "Heavy Crossbow", "Fuma Shuriken"], "ninjaTools": ["Armorsmith Kit", "Weaponsmith Kit"], "savingThrows": ["Strength", "Constitution", "Intelligence"]},
    "Cooking-Nin": {"armor": ["Light Armor", "Medium Armor"], "weapons": ["Simple Weapons"], "ninjaTools": ["Cooking Kit", "Poison Kit"], "savingThrows": ["Constitution", "Charisma", "Intelligence"]},
    "Science-Nin": {"armor": ["Light Armor", "Medium Armor"], "weapons": ["Simple Weapons", "Scientific Ninja Tools"], "ninjaTools": ["Choose 3 Ninja Tools"], "savingThrows": ["Dexterity", "Constitution", "Intelligence"]},
}

BACKGROUND_PROFICIENCY_GRANTS = {
    "Entertainer": {"skills": [fixed_skills("background-entertainer", "acrobatics", "performance")], "toolkits": ["Disguise Kit"]},
    "Genius": {"skills": [choose_skills("background-genius", 2, ["ninshou", "martial-arts", "investigation", "history"])], "toolkitChoice": True},
    "Hard Worker": {"skills": [fixed_skills("background-hard-worker", "acrobatics", "athletics")], "toolkits": ["Weaponsmith Kit"]},
    "Hermit": {"skills": [fixed_skills("background-hermit", "animal-handling", "history")], "toolkitOptions": ["Poison Kit", "Medicine Kit"]},
    "Leader": {"skills": [fixed_skills("background-leader", "history"), choose_skills("background-leader", 1, ["deception", "persuasion"])], "toolkits": ["Forgery Kit"]},
    "Noble": {"skills": [fixed_skills("background-noble", "persuasion", "insight")], "toolkits": ["Security Kit"]},
    "Student": {"skills": [fixed_skills("background-student", "acrobatics"), choose_skills("background-student", 1, ["ninshou", "investigation", "illusions", "medicine"])], "toolkits": ["Forensics Kit"]},
    "Traveler": {"skills": [fixed_skills("background-traveler", "insight", "perception")], "toolkits": ["Poison Kit"]},
    "Trouble Maker": {"skills": [fixed_skills("background-trouble-maker", "deception", "sleight-of-hand")], "toolkits": ["Forgery Kit"]},
    "Urchin": {"skills": [fixed_skills("background-urchin", "sleight-of-hand", "stealth")], "toolkits": ["Disguise Kit"]},
}


def proficiency_action(name: str, degree: int, summary: str) -> dict:
    return {"name": name, "degree": degree, "summary": summary}


PROFICIENCY_SKILLS = [
    {"id": "acrobatics", "name": "Acrobatics", "ability": "dexterity", "actions": [
        proficiency_action("Maintain Balance", 0, "Use an action and an Acrobatics check to cross unstable or difficult footing without falling prone."),
        proficiency_action("Cat Fall", 1, "Use your reaction and an Acrobatics check to lessen the impact of a fall; the DC rises with distance."),
        proficiency_action("Squeeze Through", 1, "Use an action and an Acrobatics check to pass through a space suited to a creature one size smaller."),
    ]},
    {"id": "animal-handling", "name": "Animal Handling", "ability": "wisdom", "actions": [
        proficiency_action("Calm Animal", 0, "React before violence begins to calm a hostile beast with a check against 8 + its level."),
        proficiency_action("Tame Animal", 1, "Attempt to tame a grappled, interacting, or badly wounded beast with a check against 12 + its level."),
        proficiency_action("Train Animal", 2, "Spend one downtime week and pass DC 25 to teach a tamed or bonded animal one compatible Taijutsu or Bukijutsu."),
    ]},
    {"id": "athletics", "name": "Athletics", "ability": "strength", "actions": [
        proficiency_action("Trip", 0, "Forfeit unarmed damage after a hit and win Athletics against Acrobatics to knock the target prone."),
        proficiency_action("Grapple", 0, "Use an action and win a contested Athletics check to grapple; Mastery can restrain on a result 10 higher."),
        proficiency_action("Shove", 0, "Replace an attack with a contested Athletics check to push a target; proficiency and Mastery increase distance."),
        proficiency_action("Disarm", 1, "Replace an attack with contested Athletics to knock a held item away."),
        proficiency_action("Grapple and Shove enhancements", 2, "Mastery can upgrade a decisive grapple to Restrained and doubles the proficiency-based shove distance."),
    ]},
    {"id": "chakra-control", "name": "Chakra Control", "ability": "constitution", "actions": [
        proficiency_action("Maintain Focus", 1, "Use an action and beat your jutsu save DC to avoid concentration checks until your next turn."),
        proficiency_action("Suppress Chakra", 1, "Spend one minute to conceal your chakra signature for ten minutes against passive and active detection."),
        proficiency_action("Draw Reserves", 2, "Spend ten minutes and make a DC 20 check to spend Chakra Dice and regain their results as chakra."),
    ]},
    {"id": "crafting", "name": "Crafting", "ability": "intelligence", "actions": [
        proficiency_action("Basic Repair", 0, "Use suitable smithing tools and ten minutes to remove equipment damage or penalties."),
        proficiency_action("Craft Non-Chakra Item", 1, "Create equipment using the non-chakra-enhanced crafting rules."),
        proficiency_action("Craft Armor or Weapon Seals", 1, "Create armor and weapon seals using their crafting rules."),
        proficiency_action("Identify Seals", 1, "React to inspect sealed equipment and identify its seals against a rank-based DC."),
        proficiency_action("Superior Crafting", 2, "Create higher-quality armor and weapons under the non-chakra item rules."),
        proficiency_action("Enhanced Value", 2, "Increase the market value you can produce by 150 Ryo for every downtime week spent crafting."),
    ]},
    {"id": "deception", "name": "Deception", "ability": "charisma", "actions": [
        proficiency_action("Create Diversion", 0, "Use an action against passive Perception to grant allies a d6 bonus to Stealth."),
        proficiency_action("Lie", 0, "Use an action and oppose passive Insight to make creatures believe a falsehood."),
        proficiency_action("Impersonate", 1, "Set a Deception result as the DC to see through your altered voice and mannerisms."),
        proficiency_action("Feint", 2, "Use a bonus action and beat Insight to gain a d4 on your next melee attack."),
    ]},
    {"id": "history", "name": "History", "ability": "intelligence", "actions": [
        proficiency_action("Recall Knowledge", 0, "Use an action or reaction and a History check to remember relevant facts, events, cultures, or lore."),
    ]},
    {"id": "illusions", "name": "Illusions", "ability": "wisdom", "actions": [
        proficiency_action("Cast Genjutsu Scrolls - D-Rank", 0, "Cast Genjutsu from scrolls of D-Rank or lower."),
        proficiency_action("Read the Enemy", 1, "React to Genjutsu hand seals; identify the cast and gain +2 AC or advantage on its save after a successful check."),
        proficiency_action("Identify Genjutsu", 1, "Identify an active Genjutsu and its special ending conditions against its save DC + 5."),
        proficiency_action("Cast Genjutsu Scrolls - B-Rank", 1, "Proficiency raises your Genjutsu scroll limit to B-Rank."),
        proficiency_action("Cast Genjutsu Scrolls - S-Rank", 2, "At least one Mastery rank raises your Genjutsu scroll limit to S-Rank."),
    ]},
    {"id": "insight", "name": "Insight", "ability": "wisdom", "actions": [
        proficiency_action("See Through Lies", 0, "React to a statement and oppose Deception to notice that it may be false."),
        proficiency_action("Sense Motive", 1, "Use an action to assess abnormal behavior, likely intent, and possible chakra-based influence."),
    ]},
    {"id": "intimidation", "name": "Intimidation", "ability": "charisma", "actions": [
        proficiency_action("Coerce", 0, "Spend one minute threatening a creature and contest its Wisdom save to force short-term compliance."),
        proficiency_action("Demoralize", 1, "Use an action within 30 feet and beat Insight to inflict one rank of Fear."),
        proficiency_action("Severe Demoralize", 3, "With Mastery 2, beating Demoralize by 10 or more inflicts two ranks of Fear."),
    ]},
    {"id": "investigation", "name": "Investigation", "ability": "intelligence", "actions": [
        proficiency_action("Study", 0, "Spend one minute studying a subject to uncover one relevant piece of information."),
        proficiency_action("Research", 1, "Spend ten minutes to eight hours on a deep investigation, gaining time-based bonuses to uncover important lore."),
    ]},
    {"id": "martial-arts", "name": "Martial Arts", "ability": "strength", "actions": [
        proficiency_action("Read the Enemy", 1, "React to complex movements for Taijutsu or Bukijutsu; identify the cast and gain +2 AC or advantage on its save."),
        proficiency_action("Identify Taijutsu or Bukijutsu", 1, "Identify an active physical jutsu and special ending conditions against its save DC + 5."),
    ]},
    {"id": "medicine", "name": "Medicine", "ability": "wisdom", "actions": [
        proficiency_action("Administer Aid", 0, "Use an action and pass DC 15 to remove ranks of Bleeding, Bruised, or Concussed once per rest per target."),
        proficiency_action("Treat Wounds", 1, "Spend ten minutes with a Medicine Kit to restore hit points using the target-condition healing table."),
        proficiency_action("Critical Healing", 2, "With Mastery, beating the Treat Wounds critical-success DC unlocks its larger healing result."),
    ]},
    {"id": "nature", "name": "Nature", "ability": "intelligence", "actions": []},
    {"id": "ninshou", "name": "Ninshou", "ability": "intelligence", "actions": [
        proficiency_action("Cast Ninjutsu Scrolls - D-Rank", 0, "Cast Ninjutsu from scrolls of D-Rank or lower."),
        proficiency_action("Read the Enemy", 1, "React to Ninjutsu hand seals; identify the cast and gain +2 AC or advantage on its save after a successful check."),
        proficiency_action("Identify Ninjutsu", 1, "Identify an active Ninjutsu and its special ending conditions against its save DC + 5."),
        proficiency_action("Cast Ninjutsu Scrolls - B-Rank", 1, "Proficiency raises your Ninjutsu scroll limit to B-Rank."),
        proficiency_action("Cast Ninjutsu Scrolls - S-Rank", 2, "At least one Mastery rank raises your Ninjutsu scroll limit to S-Rank."),
    ]},
    {"id": "perception", "name": "Perception", "ability": "wisdom", "actions": [
        proficiency_action("Search", 0, "Use an action and a Perception check to find something concealed or otherwise not obvious."),
    ]},
    {"id": "performance", "name": "Performance", "ability": "charisma", "actions": [
        proficiency_action("Perform - Common Audience", 0, "Perform to influence later social DCs and attract commoners, artisans, or Genin."),
        proficiency_action("Perform - Trained Audience", 1, "Proficiency lets a performance specifically attract merchants, minor nobles, or Chunin."),
        proficiency_action("Perform - Elite Audience", 2, "Mastery lets a performance specifically attract high nobility, royalty, or Jonin."),
    ]},
    {"id": "persuasion", "name": "Persuasion", "ability": "charisma", "actions": [
        proficiency_action("Gather Information", 0, "Spend two to three hours canvassing a community to learn one useful fact."),
        proficiency_action("Make an Impression or Request", 0, "Spend at least one minute building goodwill, then check against 10 + the target's Wisdom save bonus."),
        proficiency_action("Haggle", 1, "Contest a seller's passive Insight to alter an item's price."),
        proficiency_action("Upsell", 2, "Check against a GM-set DC to sell an item for 50 percent more Ryo."),
    ]},
    {"id": "sleight-of-hand", "name": "Sleight of Hand", "ability": "dexterity", "actions": [
        proficiency_action("Palm an Object", 0, "Hide a small unattended object against observers' Perception."),
        proficiency_action("Steal", 0, "Take a one-handed object unnoticed against passive Perception, with penalties for concealed items."),
        proficiency_action("Disable a Device or Trap", 1, "Use the appropriate toolkit and beat the device's Disable DC."),
        proficiency_action("Pick a Lock", 1, "Use a Security Kit under the toolkit rules to pick a lock."),
    ]},
    {"id": "stealth", "name": "Stealth", "ability": "dexterity", "actions": [
        proficiency_action("Hide", 0, "Use cover, heavy obscurement, or darkness and compare Stealth against passive Perception."),
        proficiency_action("Sneak", 0, "Move at half speed while hidden and compare Stealth against nearby creatures' passive Perception."),
    ]},
    {"id": "survival", "name": "Survival", "ability": "wisdom", "actions": [
        proficiency_action("Sense Direction", 0, "Orient yourself from natural signs; unfamiliar environments may require proficiency."),
        proficiency_action("Cover Tracks", 1, "Travel at half speed so trackers must beat your passive Survival."),
        proficiency_action("Track", 1, "Move at half travel speed and check against passive Survival, Stealth, or a GM-set DC."),
        proficiency_action("Navigate Alien Environments", 3, "Mastery 2 permits navigation through alien worlds, changing dream realms, alternate realities, or other dimensions."),
    ]},
]

SUMMON_ROLES = {
    "Striker": "This summon has Multiattack and can make two attacks with its natural or held weapons.",
    "Caster": "This summon gains additional jutsu slots equal to its rank.",
    "Controller": "This summon gains a bonus to its save DC equal to its rank bonus.",
    "Defender": "This summon gains a bonus to AC equal to its rank bonus.",
    "Lurker": "Once per turn, it deals extra d8s equal to half its level, rounded up, when it has advantage or an ally is within 5 feet of the target.",
    "Supporter": "Healing and temporary hit points gain extra dice equal to rank: D 1, C 2, B 3, A 4, S 5.",
}

SUMMON_RANKS = {
    "D-Rank": {"level": 4, "asiPoints": 0, "statMax": 20},
    "C-Rank": {"level": 8, "asiPoints": 6, "statMax": 20},
    "B-Rank": {"level": 12, "asiPoints": 12, "statMax": 22},
    "A-Rank": {"level": 16, "asiPoints": 18, "statMax": 24},
    "S-Rank": {"level": 20, "asiPoints": 24, "statMax": 26},
}

# The sheet's Lists tab contains an older summon ability matrix. The current
# summon tables in Jiraiya's Jutsu Compendium are authoritative for these core
# tribes (D-Rank base scores, before rank-based ability increases).
SUMMON_ABILITY_OVERRIDES = {
    "bear": (16, 10, 14, 10, 12, 10),
    "boar": (16, 10, 14, 10, 10, 12),
    "deer": (14, 12, 10, 10, 16, 10),
    "dogwolf": (14, 14, 14, 10, 12, 10),
    "fox": (8, 14, 12, 10, 13, 14),
    "harerabbit": (8, 14, 10, 16, 13, 11),
    "hawk": (10, 14, 12, 10, 16, 10),
    "insectswarm": (8, 13, 11, 16, 14, 10),
    "lizard": (16, 12, 12, 10, 14, 8),
    "monkey": (16, 14, 12, 10, 10, 10),
    "oxram": (16, 10, 14, 10, 12, 10),
    "rat": (8, 16, 10, 14, 13, 11),
    "shark": (16, 10, 12, 10, 10, 14),
    "slug": (10, 10, 12, 14, 16, 10),
    "snake": (10, 16, 12, 10, 10, 14),
    "spider": (14, 10, 12, 16, 10, 10),
    "tigerlion": (16, 10, 12, 10, 14, 10),
    "toad": (14, 12, 12, 14, 12, 11),
    "turtle": (13, 8, 16, 14, 11, 10),
    "weasel": (8, 16, 12, 14, 12, 10),
}

# Armor Class uses the modifier of the tribe's named Defensive Ability. These
# are taken from the current summon entries in Jiraiya's Jutsu Compendium.
# Sheet-only supplemental tribes fall back to Dexterity below so their Armor
# still responds to their allocated ability increases.
SUMMON_DEFENSIVE_ABILITIES = {
    "bear": "constitution",
    "boar": "strength",
    "deer": "dexterity",
    "dogwolf": "dexterity",
    "fox": "dexterity",
    "harerabbit": "dexterity",
    "hawk": "dexterity",
    "insectswarm": "wisdom",
    "lizard": "strength",
    "monkey": "dexterity",
    "oxram": "constitution",
    "rat": "dexterity",
    "shark": "constitution",
    "slug": "constitution",
    "snake": "dexterity",
    "spider": "strength",
    "tigerlion": "wisdom",
    "toad": "constitution",
    "turtle": "constitution",
    "weasel": "dexterity",
}


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def split_features(value) -> list[str]:
    raw = text(value)
    if not raw:
        return []
    return [item.strip() for item in re.split(r",\s*(?![^()]*\))", raw) if item.strip()]


def normalized(value: str) -> str:
    value = re.sub(r"\[[^\]]+\]", "", value)
    value = re.sub(r"\s*\([^)]*\)\s*$", "", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip().casefold()


def feature_catalog(workbook) -> dict[str, dict]:
    sheet = workbook["Feature Lookup"]
    result: dict[str, dict] = {}
    for row in sheet.iter_rows(min_row=2, max_col=7, values_only=True):
        lookup, final_name, kind, description, action, cost, data = row
        name = text(final_name) or text(lookup)
        description = text(description)
        if not name or not description:
            continue
        entry = {
            "name": name,
            "type": text(kind),
            "description": description,
            "action": text(action),
            "cost": text(cost),
            "data": text(data),
        }
        for alias in {text(lookup), name, normalized(name)}:
            if alias:
                result[normalized(alias)] = entry
    return result


def decorate_feature(name: str, catalog: dict[str, dict]) -> dict:
    entry = catalog.get(normalized(name), {})
    return {
        "name": name,
        "description": entry.get("description", ""),
        "type": entry.get("type", ""),
        "action": entry.get("action", ""),
        "cost": entry.get("cost", ""),
    }


def extract_clans(workbook, catalog) -> list[dict]:
    sheet = workbook["Lists"]
    clans = []
    for row_number in range(259, 392):
        name = text(sheet.cell(row_number, 1).value)
        if not name:
            continue
        features = {}
        for column, level in zip(range(7, 14), CLAN_LEVELS):
            names = split_features(sheet.cell(row_number, column).value)
            if names:
                features[str(level)] = [decorate_feature(item, catalog) for item in names]
        clans.append({
            "id": f"clan-{row_number}",
            "name": name,
            "speed": text(sheet.cell(row_number, 14).value),
            "featuresByLevel": features,
            "image": text(sheet.cell(row_number, 28).value),
        })
    return clans


def extract_classes(workbook, catalog) -> list[dict]:
    sheet = workbook["Lists"]
    classes = []
    for row_number in range(398, 426):
        name = text(sheet.cell(row_number, 1).value)
        if not name or name == "Example-Tester":
            continue
        features = {}
        for column, level in zip(range(7, 27), LEVELS):
            names = split_features(sheet.cell(row_number, column).value)
            if names:
                features[str(level)] = [decorate_feature(item, catalog) for item in names]
        classes.append({
            "id": f"class-{row_number}",
            "name": name,
            "hitDie": text(sheet.cell(row_number, 27).value),
            "chakraDie": text(sheet.cell(row_number, 28).value),
            "jutsuKnown": text(sheet.cell(row_number, 40).value),
            "startingEquipment": [
                {"id": f"{row_number}-{index}", "choices": choices}
                for index, choices in enumerate(STARTING_EQUIPMENT.get(name, []), 1)
            ],
            "skillGrants": CLASS_PROFICIENCY_GRANTS.get(name, []),
            "combatProficiencies": CLASS_COMBAT_PROFICIENCIES.get(name),
            "featuresByLevel": features,
        })
    return classes


def extract_subclasses(workbook, catalog) -> list[dict]:
    sheet = workbook["Lists"]
    subclasses = []
    for row_number in range(432, 632):
        class_name = text(sheet.cell(row_number, 1).value)
        name = text(sheet.cell(row_number, 7).value)
        if not class_name or not name:
            continue
        features = {}
        for column, level in zip(range(16, 36), LEVELS):
            names = split_features(sheet.cell(row_number, column).value)
            if names:
                features[str(level)] = [decorate_feature(item, catalog) for item in names]
        subclasses.append({
            "id": f"subclass-{row_number}",
            "className": class_name,
            "name": name,
            "featuresByLevel": features,
        })
    return subclasses


def extract_backgrounds(workbook, catalog) -> list[dict]:
    sheet = workbook["Lists"]
    backgrounds = []
    for row_number in range(645, 720):
        name = text(sheet.cell(row_number, 1).value)
        if not name:
            continue
        feature_name = text(sheet.cell(row_number, 7).value)
        backgrounds.append({
            "id": f"background-{row_number}",
            "name": name,
            "feature": decorate_feature(feature_name, catalog) if feature_name else None,
            "proficiencyGrants": BACKGROUND_PROFICIENCY_GRANTS.get(name),
        })
    return backgrounds


def extract_jutsu(workbook) -> list[dict]:
    sheet = workbook["Jutsu List (Courtesy of Kane)"]
    jutsu = []
    for row_number in range(2, 2001):
        values = [text(sheet.cell(row_number, column).value) for column in range(1, 12)]
        name, kind, rank, casting_time, range_, duration, components, cost, keywords, description, higher = values
        if not name or not kind or name in {"Feat Titles", "Jutsu Names"}:
            continue
        if rank and "Rank" not in rank:
            continue
        jutsu.append({
            "id": f"jutsu-{row_number}",
            "name": name,
            "classification": kind,
            "rank": rank,
            "castingTime": casting_time,
            "range": range_,
            "duration": duration,
            "components": components,
            "cost": cost,
            "keywords": keywords,
            "description": description,
            "higherRanks": higher,
        })
    return jutsu


def extract_equipment(workbook) -> dict:
    sheet = workbook["Lists"]
    armor = []
    for row_number in range(8, 49):
        name = text(sheet.cell(row_number, 1).value)
        if not name or name == "Skin" or name.startswith("[["):
            continue
        armor.append({
            "id": f"armor-{row_number}", "name": name, "type": text(sheet.cell(row_number, 7).value),
            "armorBonus": sheet.cell(row_number, 25).value or 0, "maxModifier": sheet.cell(row_number, 21).value,
            "bulk": sheet.cell(row_number, 33).value or 0, "price": ARMOR_PRICES.get(name.casefold()),
        })
    weapons = []
    for row_number in range(61, 167):
        name = text(sheet.cell(row_number, 1).value)
        if not name or name.startswith("[["):
            continue
        market_name = re.sub(r"\s*\((?:Two Hands)\)\s*$", "", name, flags=re.I).casefold()
        price, category = WEAPON_MARKET.get(market_name, (None, ""))
        weapons.append({
            "id": f"weapon-{row_number}", "name": name, "damage": text(sheet.cell(row_number, 8).value),
            "ability": text(sheet.cell(row_number, 11).value), "damageType": text(sheet.cell(row_number, 25).value),
            "properties": text(sheet.cell(row_number, 35).value), "price": price, "category": category,
        })
    return {"armor": armor, "weapons": weapons}


def summon_key(value: str) -> str:
    value = value.casefold().replace("predator birds", "").replace("predator bird", "").replace("primate", "")
    value = value.replace("rabbits", "rabbit")
    return re.sub(r"[^a-z0-9]", "", value)


def extract_summons(workbook) -> dict:
    sheet = workbook["Lists"]
    summon_sheet = workbook["Summon Sheet"]
    trait_descriptions = {}
    for row_number in range(3, 1003):
        trait_name = text(summon_sheet.cell(row_number, 167).value)
        description = text(summon_sheet.cell(row_number, 170).value)
        if trait_name and description:
            trait_descriptions[trait_name.casefold()] = description
    rank_columns = {
        "D-Rank": {"size": 91, "slots": 80, "speed": 112, "features": (141, 142, 143)},
        "C-Rank": {"size": 92, "slots": 81, "speed": 113, "features": (146, 147, 148)},
        "B-Rank": {"size": 93, "slots": 82, "speed": 114, "features": (151, 152, 153)},
        "A-Rank": {"size": 94, "slots": 83, "speed": 115, "features": (156, 157, 158)},
        "S-Rank": {"size": 95, "slots": 84, "speed": 116, "features": (161, 162, 163)},
    }
    advancement = {}
    for row_number in range(4, 74):
        name = text(summon_sheet.cell(row_number, 65).value)
        if not name:
            continue
        ranks = {}
        for rank, details in rank_columns.items():
            features = []
            for column in details["features"]:
                feature_name = text(summon_sheet.cell(row_number, column).value)
                if feature_name and feature_name != "-":
                    features.append({
                        "name": feature_name,
                        "description": trait_descriptions.get(feature_name.casefold(), ""),
                    })
            ranks[rank] = {
                **SUMMON_RANKS[rank],
                "sizeRange": text(summon_sheet.cell(row_number, details["size"]).value),
                "jutsuSlots": summon_sheet.cell(row_number, details["slots"]).value or 0,
                "speed": summon_sheet.cell(row_number, details["speed"]).value or 0,
                "features": features,
            }
        advancement[summon_key(name)] = {
            "summonType": text(summon_sheet.cell(row_number, 70).value),
            "toughness": summon_sheet.cell(row_number, 75).value or 0,
            "castingModifier": summon_sheet.cell(row_number, 96).value or 0,
            "defensiveModifier": summon_sheet.cell(row_number, 101).value or 0,
            "ranks": ranks,
        }
    tribes = []
    for row_number in range(173, 209):
        name = text(sheet.cell(row_number, 1).value)
        if not name:
            continue
        progression = advancement.get(summon_key(name), {})
        sheet_abilities = tuple(sheet.cell(row_number, column).value or 10 for column in (39, 41, 43, 45, 47, 49))
        ability_scores = SUMMON_ABILITY_OVERRIDES.get(summon_key(name), sheet_abilities)
        defensive_ability = SUMMON_DEFENSIVE_ABILITIES.get(summon_key(name), "dexterity")
        tribes.append({
            "id": f"tribe-{row_number}", "name": name,
            "savingThrows": [text(sheet.cell(row_number, 7).value), text(sheet.cell(row_number, 9).value)],
            "skills": text(sheet.cell(row_number, 11).value), "senses": text(sheet.cell(row_number, 15).value),
            "jutsuAccess": text(sheet.cell(row_number, 19).value), "attackAbility": text(sheet.cell(row_number, 25).value),
            "naturalWeapons": text(sheet.cell(row_number, 29).value),
            "abilities": {
                "strength": ability_scores[0], "dexterity": ability_scores[1],
                "constitution": ability_scores[2], "intelligence": ability_scores[3],
                "wisdom": ability_scores[4], "charisma": ability_scores[5],
            },
            "defensiveAbility": defensive_ability,
            "hitDie": sheet.cell(row_number, 51).value, "chakraDie": sheet.cell(row_number, 53).value,
            **progression,
        })
    templates = []
    for row_number in range(215, 253):
        tribe = text(sheet.cell(row_number, 13).value)
        if not tribe:
            continue
        templates.append({
            "id": f"summon-{row_number}", "name": text(sheet.cell(row_number, 1).value),
            "rank": text(sheet.cell(row_number, 9).value), "tribe": tribe,
            "features": [text(sheet.cell(row_number, col).value) for col in (19, 27, 35, 43) if text(sheet.cell(row_number, col).value) not in ("", "-")],
            "speed": sheet.cell(row_number, 51).value, "size": text(sheet.cell(row_number, 53).value),
        })
    return {
        "tribes": tribes,
        "templates": templates,
        "roles": [{"name": name, "description": description} for name, description in SUMMON_ROLES.items()],
        "ranks": SUMMON_RANKS,
    }


def pdf_sources(paths: list[Path]) -> list[dict]:
    sources = []
    for path in paths:
        if not path.exists():
            continue
        reader = PdfReader(str(path))
        sources.append({"name": path.name, "pages": len(reader.pages)})
    return sources


def main() -> None:
    if len(sys.argv) < 3:
        raise SystemExit("Expected input workbook, output JSON, and optional PDF paths")
    workbook_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    pdf_paths = [Path(value) for value in sys.argv[3:]]
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    catalog = feature_catalog(workbook)
    payload = {
        "schemaVersion": 1,
        "source": {
            "sheetTitle": "Copy of N5E Character Sheet - Final Version (Make a Copy)",
            "sheetVersion": text(workbook["Character Info"]["X3"].value),
            "pdfs": pdf_sources(pdf_paths),
        },
        "classes": extract_classes(workbook, catalog),
        "subclasses": extract_subclasses(workbook, catalog),
        "clans": extract_clans(workbook, catalog),
        "backgrounds": extract_backgrounds(workbook, catalog),
        "jutsu": extract_jutsu(workbook),
        "equipment": extract_equipment(workbook),
        "summons": extract_summons(workbook),
        "proficiencies": {
            "degrees": PROFICIENCY_DEGREES,
            "skills": PROFICIENCY_SKILLS,
            "masteryCaps": [
                {"minLevel": 1, "maxLevel": 6, "rank": 1},
                {"minLevel": 7, "maxLevel": 11, "rank": 2},
                {"minLevel": 12, "maxLevel": 20, "rank": 3},
            ],
        },
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    counts = {key: len(payload[key]) for key in ("classes", "subclasses", "clans", "backgrounds", "jutsu")}
    print(json.dumps({"output": str(output_path), "counts": counts}, indent=2))


if __name__ == "__main__":
    main()
